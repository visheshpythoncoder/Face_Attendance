import tkinter as tk
from tkinter import messagebox
import openpyxl
import os

EXCEL_PATH = "user_data.xlsx"

def create_excel_file():
    if not os.path.exists(EXCEL_PATH):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "LoginData"
        sheet.append(["Username", "Password", "FaceImagePath"])
        workbook.save(EXCEL_PATH)

def login():
    username = username_entry.get()
    password = password_entry.get()

    if not username or not password:
        messagebox.showerror("Error", "Please fill in all fields.")
        return

    workbook = openpyxl.load_workbook(EXCEL_PATH)
    sheet = workbook["LoginData"]

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) < 3:
            continue  # Skip rows with missing data
        stored_username, stored_password, stored_face_image_path = row
        if username == stored_username and password == stored_password:
            messagebox.showinfo("Login Success", f"Welcome {username}!\nFace path: {stored_face_image_path}")
            return

    messagebox.showerror("Login Failed", "Invalid username or password.")

def register():
    username = username_entry.get()
    password = password_entry.get()
    face_image_path = face_path_entry.get()

    if not username or not password or not face_image_path:
        messagebox.showerror("Error", "Please fill in all fields.")
        return

    workbook = openpyxl.load_workbook(EXCEL_PATH)
    sheet = workbook["LoginData"]

    # Check if username already exists
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == username:
            messagebox.showerror("Error", "Username already exists.")
            return

    sheet.append([username, password, face_image_path])
    workbook.save(EXCEL_PATH)
    messagebox.showinfo("Success", "Registration successful.")

# ---------------------- UI Setup ----------------------
create_excel_file()

root = tk.Tk()
root.title("Login System")
root.geometry("400x300")

# Configure grid weights
root.grid_rowconfigure(0, weight=1)
root.grid_rowconfigure(6, weight=1)
root.grid_columnconfigure(0, weight=1)
root.grid_columnconfigure(1, weight=1)

tk.Label(root, text="Username:").grid(row=1, column=0, sticky="e", padx=10, pady=5)
username_entry = tk.Entry(root)
username_entry.grid(row=1, column=1, padx=10, pady=5)

tk.Label(root, text="Password:").grid(row=2, column=0, sticky="e", padx=10, pady=5)
password_entry = tk.Entry(root, show="*")
password_entry.grid(row=2, column=1, padx=10, pady=5)

tk.Label(root, text="Face Image Path:").grid(row=3, column=0, sticky="e", padx=10, pady=5)
face_path_entry = tk.Entry(root)
face_path_entry.grid(row=3, column=1, padx=10, pady=5)

tk.Button(root, text="Login", command=login).grid(row=4, column=0, pady=20)
tk.Button(root, text="Register", command=register).grid(row=4, column=1, pady=20)

root.mainloop()
